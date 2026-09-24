import datetime
import os
import pandas as pd  # type: ignore
import openpyxl
from openpyxl.utils import get_column_letter
from utils.logs_escrita import log_info, log_error  # type: ignore


def atualizar_chamados_zero_tabela(caminho_work, caminho_zero):
    try:
        log_info(f"Lendo dados de origem: {caminho_work}")
        df_work = pd.read_excel(caminho_work, engine="openpyxl")

        if df_work.empty:
            log_info("Arquivo de origem vazio. Nenhuma ação realizada.")
            return

        # 2. Cria a coluna 'Atualização Arquivo' com Data e Hora atual
        data_hora_atual = datetime.datetime.now()
        df_work['Atualização Arquivo'] = data_hora_atual

        # 3. Mapeia a ordem exata das colunas de destino
        colunas_destino = [
            'CODIGO_TRATADO', 'NOME', 'DISPONIBILIDADO_ABONO', 'DATA_REPORT', 'VERSAO',
            'VIP', 'RESPONSAVEL', 'Atualização Arquivo'
        ]

        # Garante que todas as colunas necessárias existam
        for col in colunas_destino:
            if col not in df_work.columns:
                df_work[col] = None

        df_final = df_work[colunas_destino].copy()

        # 4. Formatações específicas por coluna para evitar o '00:00:00' indesejado
        if 'DATA_REPORT' in df_final.columns:
            # df_final['DATA_REPORT'] = pd.to_datetime(df_final['DATA_REPORT']).dt.strftime('%d/%m/%Y')
            df_final['DATA_REPORT'] = pd.to_datetime(df_final['DATA_REPORT']).dt.date
            
        if 'Atualização Arquivo' in df_final.columns:
            df_final['Atualização Arquivo'] = pd.to_datetime(df_final['Atualização Arquivo']).dt.strftime('%d/%m/%Y %H:%M:%S')

        # Substitui valores nulos por None para o openpyxl inserir células vazias corretamente
        df_final = df_final.where(pd.notnull(df_final), None)

        # 5. Processamento via openpyxl na planilha de destino
        log_info(f"Abrindo pasta de trabalho: {caminho_zero}")
        wb = openpyxl.load_workbook(caminho_zero)
        ws = wb.active
        ws = wb["Sheet1"]  # Substitua "Sheet1" pelo nome correto da aba, se necessário

        # Deleta da linha 2 até a última linha com dados (mantendo o cabeçalho na linha 1)
        if ws.max_row > 1:
            ws.delete_rows(2, amount=ws.max_row - 1)

        # Escreve as novas linhas a partir da linha 2
        for row_data in df_final.values.tolist():
            ws.append(row_data)

        # 6. Ajusta dinamicamente o intervalo da Tabela do Excel
        total_linhas = len(df_final) + 1  # +1 para incluir a linha de cabeçalho
        total_colunas = len(colunas_destino)
        ultima_coluna_letra = get_column_letter(total_colunas)

        if ws.tables:
            for tabela in ws.tables.values():
                tabela.ref = f"A1:{ultima_coluna_letra}{total_linhas}"

        # Salva o arquivo preservando a Tabela do Excel e sua formatação visual
        wb.save(caminho_zero)
        log_info(f"Tabela em {caminho_zero} zerada e recarregada com sucesso!")

    except Exception as e:
        log_error(f"Erro ao atualizar a tabela do arquivo DISP_ATUALIZACAO: {str(e)}")


def gat_manual():
    from utils.caminhos import (
        PASTA_ATUALIZA_DISP, PASTA_TRATAR, PASTABKP_DESTINO
    )
    
    ARQUIVO_WORK = os.path.join(PASTA_TRATAR, "analise_manual.xlsx")
    
    # Verifica diretamente se o arquivo exato existe na pasta
    if os.path.exists(ARQUIVO_WORK):
        ARQUIVO_DISP = PASTA_ATUALIZA_DISP            
        atualizar_chamados_zero_tabela(ARQUIVO_WORK, ARQUIVO_DISP)
        
        os.system(f'move "{ARQUIVO_WORK}" "{PASTABKP_DESTINO}"')
    else:
        log_info("Nenhum arquivo 'analise_manual.xlsx' encontrado na pasta de origem.")
