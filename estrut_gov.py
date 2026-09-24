import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ARQUIVO_ORIGEM = r"C:\Users\mayzon.santos\Music\cam1\TESTE02.xlsx"
ARQUIVO_DESTINO = r"C:\Users\mayzon.santos\Music\cam1\TESTE01.xlsx"

ABA_ORIGEM = "arq1"
ABA_DESTINO = "BASE"

# Ler origem

df_origem = pd.read_excel(
    ARQUIVO_ORIGEM,
    sheet_name=ABA_ORIGEM
)

print(f"Registros encontrados: {len(df_origem)}")

# Abrir destino

wb = load_workbook(ARQUIVO_DESTINO)
ws = wb[ABA_DESTINO]

# Localizar a tabela da planilha

if not ws.tables:
    raise Exception(
        f"Nenhuma tabela encontrada na aba '{ABA_DESTINO}'."
    )

# Pega a primeira tabela da aba
tabela = list(ws.tables.values())[0]

print(f"Tabela encontrada: {tabela.name}")
print(f"Intervalo atual da tabela: {tabela.ref}")

#  Descobrir o intervalo atual da tabela


from openpyxl.utils.cell import range_boundaries

min_col, min_row, max_col, max_row = range_boundaries(
    tabela.ref
)

print(f"Última linha da tabela: {max_row}")

# Inserir os dados após a última linha da tabela


proxima_linha = max_row + 1

for linha in df_origem.itertuples(index=False, name=None):

    for coluna, valor in enumerate(linha, start=min_col):

        ws.cell(
            row=proxima_linha,
            column=coluna,
            value=valor
        )

    proxima_linha += 1

ultima_linha = proxima_linha - 1

# Expandir a tabela


nova_referencia = (
    f"{get_column_letter(min_col)}{min_row}:"
    f"{get_column_letter(max_col)}{ultima_linha}"
)

tabela.ref = nova_referencia

print(f"Nova referência da tabela: {tabela.ref}")

# Salvar

wb.save(ARQUIVO_DESTINO)


print()
print("======================================")
print("PROCESSO CONCLUÍDO")
print("======================================")
print(f"Registros adicionados: {len(df_origem)}")
print(f"Primeira linha:       {max_row + 1}")
print(f"Última linha:         {ultima_linha}")
print(f"Tabela:               {tabela.name}")
print(f"Novo intervalo:       {tabela.ref}")
print("======================================")


# # ARQUIVO EXCEL CELULAR
# import pandas as pd
# from openpyxl import load_workbook

# ARQUIVO_ORIGEM = r"C:\Users\mayzon.santos\Music\cam1\TESTE02.xlsx"
# ARQUIVO_DESTINO = r"C:\Users\mayzon.santos\Music\cam1\TESTE01.xlsx"

# ABA_ORIGEM = "arq1"
# ABA_DESTINO = "BASE"

# df_origem = pd.read_excel(
#     ARQUIVO_ORIGEM,
#     sheet_name=ABA_ORIGEM
# )

# print(f"Registros encontrados na origem: {len(df_origem)}")

# wb = load_workbook(ARQUIVO_DESTINO)
# ws = wb[ABA_DESTINO]

# ultima_linha = 0

# for linha in range(ws.max_row, 0, -1):

#     if any(
#         ws.cell(linha, coluna).value is not None
#         for coluna in range(1, 11)
#     ):
#         ultima_linha = linha
#         break

# print(f"Última linha atual da BASE: {ultima_linha}")

# proxima_linha = ultima_linha + 1

# for linha in df_origem.itertuples(index=False, name=None):

#     for coluna, valor in enumerate(linha, start=1):

#         ws.cell(
#             row=proxima_linha,
#             column=coluna,
#             value=valor
#         )

#     proxima_linha += 1


# ultima_linha_inserida = proxima_linha - 1

# wb.save(ARQUIVO_DESTINO)

# print()
# print("======================================")
# print("PROCESSO CONCLUÍDO")
# print("======================================")
# print(f"Registros adicionados : {len(df_origem)}")
# print(f"Primeira linha        : {ultima_linha + 1}")
# print(f"Última linha          : {ultima_linha_inserida}")
# print(f"Arquivo               : {ARQUIVO_DESTINO}")
# print("======================================")


import os
# import pandas as pd
# import win32com.client as win32


# ARQUIVO_ORIGEM = r"C:\Users\mayzon.santos\Music\cam1\TESTE02.xlsx"
# ARQUIVO_DESTINO = r"C:\Users\mayzon.santos\Music\cam1\TESTE01.xlsx"

# ABA_ORIGEM = "arq1"
# ABA_DESTINO = "BASE"
# NOME_TABELA = "Tabela1"


# # ============================================================
# # Função para converter valores do Pandas para o Excel
# # ============================================================

# def preparar_valor(valor):

#     # Valor vazio
#     if pd.isna(valor):
#         return None

#     # Timestamp do Pandas
#     if isinstance(valor, pd.Timestamp):

#         # Se possuir timezone, remove o timezone
#         if valor.tzinfo is not None:
#             valor = valor.tz_localize(None)

#         return valor.to_pydatetime()

#     # Outros tipos de data
#     if isinstance(valor, pd.DatetimeIndex):
#         return valor.to_pydatetime()

#     return valor


# # ============================================================
# # 1. Ler dados da origem
# # ============================================================

# df = pd.read_excel(
#     ARQUIVO_ORIGEM,
#     sheet_name=ABA_ORIGEM
# )

# print(f"Registros encontrados: {len(df)}")
# print(f"Colunas encontradas: {len(df.columns)}")


# # ============================================================
# # 2. Preparar dados para o Excel
# # ============================================================

# dados = []

# for linha in df.iloc[:, :10].itertuples(
#     index=False,
#     name=None
# ):

#     nova_linha = tuple(
#         preparar_valor(valor)
#         for valor in linha
#     )

#     dados.append(nova_linha)


# # ============================================================
# # 3. Abrir Excel
# # ============================================================

# excel = win32.DispatchEx("Excel.Application")

# excel.Visible = False
# excel.DisplayAlerts = False
# excel.ScreenUpdating = False

# wb = None

# try:

#     # ========================================================
#     # 4. Abrir arquivo destino
#     # ========================================================

#     wb = excel.Workbooks.Open(
#         os.path.abspath(ARQUIVO_DESTINO)
#     )

#     ws = wb.Worksheets(ABA_DESTINO)

#     print(f"Aba encontrada: {ws.Name}")


#     # ========================================================
#     # 5. Localizar tabela
#     # ========================================================

#     tabela = ws.ListObjects(NOME_TABELA)

#     print(f"Tabela encontrada: {tabela.Name}")


#     # ========================================================
#     # 6. Informações atuais da tabela
#     # ========================================================

#     intervalo_atual = tabela.Range

#     primeira_linha = intervalo_atual.Row

#     ultima_linha = (
#         intervalo_atual.Row
#         + intervalo_atual.Rows.Count
#         - 1
#     )

#     primeira_coluna = intervalo_atual.Column

#     quantidade_colunas = (
#         intervalo_atual.Columns.Count
#     )

#     print(f"Primeira linha da tabela: {primeira_linha}")
#     print(f"Última linha da tabela: {ultima_linha}")
#     print(f"Quantidade de colunas: {quantidade_colunas}")


#     # ========================================================
#     # 7. Quantidade de registros novos
#     # ========================================================

#     quantidade_registros = len(dados)

#     if quantidade_registros == 0:

#         print("Nenhum registro para adicionar.")

#     else:

#         primeira_nova_linha = ultima_linha + 1

#         nova_ultima_linha = (
#             ultima_linha
#             + quantidade_registros
#         )


#         # ====================================================
#         # 8. Expandir a tabela
#         # ====================================================

#         nova_range = ws.Range(
#             ws.Cells(
#                 primeira_linha,
#                 primeira_coluna
#             ),
#             ws.Cells(
#                 nova_ultima_linha,
#                 primeira_coluna
#                 + quantidade_colunas
#                 - 1
#             )
#         )

#         tabela.Resize(nova_range)

#         print(
#             f"Tabela expandida até a linha "
#             f"{nova_ultima_linha}"
#         )


#         # ====================================================
#         # 9. Inserir A:J
#         # ====================================================

#         intervalo_dados = ws.Range(
#             ws.Cells(
#                 primeira_nova_linha,
#                 1
#             ),
#             ws.Cells(
#                 nova_ultima_linha,
#                 10
#             )
#         )

#         intervalo_dados.Value = dados

#         print("Dados A:J inseridos.")


#         # ====================================================
#         # 10. Replicar fórmulas K:M
#         # ====================================================

#         print("Replicando fórmulas K:M...")

#         for coluna in range(11, 14):

#             origem = ws.Cells(
#                 ultima_linha,
#                 coluna
#             )

#             destino = ws.Range(
#                 ws.Cells(
#                     ultima_linha,
#                     coluna
#                 ),
#                 ws.Cells(
#                     nova_ultima_linha,
#                     coluna
#                 )
#             )

#             destino.FillDown()

#         print("Fórmulas K:M replicadas.")


#         # ====================================================
#         # 11. Forçar cálculo
#         # ====================================================

#         excel.CalculateFull()


#         # ====================================================
#         # 12. Salvar
#         # ====================================================

#         wb.Save()


#         # ====================================================
#         # 13. Resultado
#         # ====================================================

#         print()
#         print("==========================================")
#         print("PROCESSO CONCLUÍDO")
#         print("==========================================")
#         print(
#             f"Registros adicionados : "
#             f"{quantidade_registros}"
#         )
#         print(
#             f"Primeira nova linha   : "
#             f"{primeira_nova_linha}"
#         )
#         print(
#             f"Última nova linha     : "
#             f"{nova_ultima_linha}"
#         )
#         print(
#             f"Tabela                : "
#             f"{tabela.Name}"
#         )
#         print(
#             f"Intervalo final       : "
#             f"{tabela.Range.Address}"
#         )
#         print("A:J                   : Dados")
#         print("K:M                   : Fórmulas")
#         print("Tabela expandida      : SIM")
#         print("==========================================")


# finally:

#     if wb is not None:

#         try:
#             wb.Close(SaveChanges=False)
#         except:
#             pass

#     try:
#         excel.Quit()
#     except:
#         pass
