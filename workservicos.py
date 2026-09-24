import pandas as pd  # type: ignore
import datetime
import openpyxl
import os
from utils.logs_escrita import log_info, log_error  # type: ignore


def atualizar_chamados_zero_tabela(caminho_work, caminho_zero):
    try:
        log_info(f"Lendo dados de origem: {caminho_work}")
        df_work = pd.read_excel(caminho_work, engine="openpyxl")

        if df_work.empty:
            log_info("Arquivo de origem vazio. Nenhuma ação realizada.")
            return

        # 1. Tratamento da coluna 'Cód' (extraindo o número antes do hífen)
        if 'Código' in df_work.columns:
            df_work['Cód'] = (
                df_work['Código']
                .astype(str)
                .str.split('-')
                .str[0]
                .str.strip()
                .str.replace(r'\.0$', '', regex=True)
            )
        else:
            df_work['Cód'] = None

        # 2. Cria a coluna 'Atualização Arquivo' com Data e Hora atual
        data_hora_atual = datetime.datetime.now().strftime("%d/%m/%Y %H:%M:%S")
        df_work['Atualização Arquivo'] = data_hora_atual

        # 3. Mapeia a ordem exata das 31 colunas esperadas na CHAMADOS_ZERO
        colunas_destino = [
            'Cód', 'Elemento Primário', 'Código', 'Empresa', 'Número',
            'Grupo designado', 'Tipo de Serviço', 'Estado', 'Subestado',
            'Atribuído a', 'Descrição resumida', 'Integration ID', 'CI Impactado',
            'Criação em', 'Aberto por', 'Atualizado(a) em', 'Atualização de',
            'Anotações de trabalho', 'Encerrado', 'Encerrado por', 'CNPJ',
            'Descrição', 'Serviço de Negócio', 'Email para contato', 'Causa do erro',
            'Modelo', 'Nome para contato', 'Telefone para contato',
            'Início de viagem programado', 'Início programado', 'Atualização Arquivo'
        ]

        # Garante que todas as colunas necessárias existam
        for col in colunas_destino:
            if col not in df_work.columns:
                df_work[col] = None

        df_final = df_work[colunas_destino]

        # Função de formatação tratando NaT e datas nulas em primeiro lugar
        def formatar_valores(val):
            if pd.isna(val) or val is pd.NaT:
                return None
            if isinstance(val, (pd.Timestamp, datetime.datetime)):
                return val.strftime("%d/%m/%Y %H:%M:%S")
            return val

        # Aplica a formatação em cada elemento
        if hasattr(df_final, 'map'):
            df_final = df_final.map(formatar_valores)
        else:
            df_final = df_final.applymap(formatar_valores)

        # 4. Processamento via openpyxl na CHAMADOS_ZERO.xlsx
        log_info(f"Abrindo pasta de trabalho: {caminho_zero}")
        wb = openpyxl.load_workbook(caminho_zero)
        ws = wb.active  # Ou wb["BASE_CHAMADOS"]

        # Deleta da linha 2 até a última linha com dados
        if ws.max_row > 1:
            ws.delete_rows(2, amount=ws.max_row - 1)

        # Escreve as novas linhas a partir da linha 2
        for row_data in df_final.values.tolist():
            ws.append(row_data)

        # 5. Ajusta o intervalo da Tabela do Excel usando ws.tables.values()
        total_linhas = len(df_final) + 1  # 1 para o cabeçalho
        if ws.tables:
            for tabela in ws.tables.values():
                tabela.ref = f"A1:AE{total_linhas}"

        # Salva o arquivo preservando a Tabela do Excel
        wb.save(caminho_zero)
        log_info(f"Tabela em {caminho_zero} zerada e recarregada com sucesso!")

    except Exception as e:
        log_error(f"Erro ao atualizar a tabela do arquivo CHAMADOS_ZERO: {str(e)}")


# # Chamada do código
# if __name__ == "__main__":
#     from utils.caminhos import (
#     PASTA_CHAMADOS_ZERO, PASTA_TRATAR
#     )

#     ARQUIVO_WORK = os.path.join(PASTA_TRATAR, "1 - #Work - Serviços.xlsx")
#     ARQUIVO_ZERO = os.path.join(PASTA_CHAMADOS_ZERO, "CHAMADOS_ZERO.xlsx")

#     atualizar_chamados_zero_tabela(ARQUIVO_WORK, ARQUIVO_ZERO)

def gat_servicos():
    from utils.caminhos import (
        PASTA_CHAMADOS_ZERO, PASTA_TRATAR, PASTABKP_DESTINO
    )
    
    ARQUIVO_WORK = os.path.join(PASTA_TRATAR, "1 - #Work - Serviços.xlsx")
    
    # Verifica diretamente se o arquivo exato existe na pasta
    if os.path.exists(ARQUIVO_WORK):          
        atualizar_chamados_zero_tabela(ARQUIVO_WORK, PASTA_CHAMADOS_ZERO)
        
        os.system(f'move "{ARQUIVO_WORK}" "{PASTABKP_DESTINO}"')
    else:
        log_info("Nenhum arquivo '1 - #Work - Serviços.xlsx' encontrado na pasta de origem.")
